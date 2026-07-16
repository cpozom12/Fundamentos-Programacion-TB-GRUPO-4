from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("TF-GRUPO4_v2_CORREGIDO.pptx")

NAVY = RGBColor(25, 34, 83)
BLUE = RGBColor(72, 83, 238)
BLUE_DARK = RGBColor(50, 60, 170)
BLUE_LIGHT = RGBColor(220, 225, 255)
BG = RGBColor(242, 243, 249)
WHITE = RGBColor(255, 255, 255)
TEXT = RGBColor(31, 39, 85)
MUTED = RGBColor(91, 97, 127)
GREEN = RGBColor(37, 171, 125)
ORANGE = RGBColor(244, 151, 56)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line else fill
    return shape


def text(slide, value, x, y, w, h, size=22, bold=False, color=TEXT,
         align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=19, color=TEXT, spacing=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.line_spacing = 1.05
    return box


def base_slide(title, number=None):
    slide = prs.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, BG)
    rect(slide, 10.65, 0, 2.683, 7.5, NAVY)
    rect(slide, 10.35, 0, 0.3, 7.5, BLUE)
    rect(slide, 0, 0, 1.2, 0.22, BLUE)
    rect(slide, 0, 0.22, 0.8, 0.18, BLUE_DARK)
    rect(slide, 0, 7.12, 1.0, 0.38, BLUE)
    text(slide, "EPE | UPC", 0.8, 0.65, 2.1, 0.35, 18, True, NAVY)
    text(slide, title, 0.85, 1.45, 8.9, 0.9, 27, True, NAVY)
    if number:
        text(slide, f"{number:02d}", 11.45, 6.45, 1.1, 0.45, 18, True, WHITE, PP_ALIGN.CENTER)
    return slide


def badge(slide, label, x, y, w=2.0, fill=BLUE, color=WHITE):
    rect(slide, x, y, w, 0.48, fill, True)
    text(slide, label, x, y + 0.04, w, 0.32, 13, True, color,
         PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def flow_box(slide, label, x, y, w=2.0, h=0.8, fill=WHITE):
    rect(slide, x, y, w, h, fill, True, BLUE_LIGHT)
    text(slide, label, x + 0.12, y + 0.12, w - 0.24, h - 0.2, 15,
         True, NAVY, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


# 1. Portada
s = prs.slides.add_slide(blank)
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 0, 0, 2.0, 7.5, BLUE_DARK)
rect(s, 1.3, 0, 0.35, 7.5, BLUE)
text(s, "UPC", 5.65, 0.75, 2.0, 0.55, 28, True, WHITE, PP_ALIGN.CENTER)
text(s, "FUNDAMENTOS DE\nPROGRAMACIÓN", 2.2, 2.0, 8.9, 1.55, 39,
     True, WHITE, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(s, "LIQUIDADOR DE HORAS EXTRAS AUTOMATIZADO", 2.5, 3.75, 8.3,
     0.45, 18, False, BLUE_LIGHT, PP_ALIGN.CENTER)
badge(s, "TRABAJO FINAL - FP1 / GRUPO 4", 4.3, 4.7, 4.8, BLUE)

# 2. Contenido
s = base_slide("CONTENIDO", 2)
left = ["El problema", "Objetivo del proyecto", "Alternativas evaluadas",
        "Solución elegida", "Detalle del proceso"]
right = ["Algoritmo (pseudocódigo)", "Herramientas y versiones",
         "Demostración en vivo", "Resultados", "Conclusiones"]
for i, item in enumerate(left, 1):
    badge(s, f"{i:02d}", 1.0, 2.45 + (i - 1) * 0.65, 0.65, WHITE, BLUE)
    text(s, item.upper(), 1.8, 2.5 + (i - 1) * 0.65, 3.8, 0.35, 16, True, TEXT)
for i, item in enumerate(right, 6):
    badge(s, f"{i:02d}", 5.95, 2.45 + (i - 6) * 0.65, 0.65, WHITE, BLUE)
    text(s, item.upper(), 6.75, 2.5 + (i - 6) * 0.65, 3.4, 0.35, 16, True, TEXT)

# 3. El problema
s = base_slide("EL PROBLEMA", 3)
bullets(s, [
    "El pago semanal de horas extra se calcula manualmente.",
    "El supervisor utiliza una calculadora básica y apuntes físicos.",
    "El recargo escalonado de 25% y 35% genera errores de cálculo y redondeo.",
    "La consecuencia son reclamos, reprocesos y riesgo ante una fiscalización de SUNAFIL."
], 0.95, 2.45, 7.8, 3.4, 19)
rect(s, 9.0, 2.1, 1.25, 3.7, WHITE, True)
text(s, "!", 9.22, 2.45, 0.8, 1.0, 58, True, ORANGE, PP_ALIGN.CENTER)
text(s, "Cálculo\nmanual", 9.1, 3.6, 1.05, 1.1, 17, True, NAVY, PP_ALIGN.CENTER)

# 4. Objetivo
s = base_slide("OBJETIVO DEL PROYECTO", 4)
text(s, "Objetivo general", 0.95, 2.35, 3.0, 0.4, 20, True, BLUE)
text(s, "Automatizar el cálculo de horas extra aplicando el recargo de ley (25%/35%) para reducir errores manuales.",
     0.95, 2.85, 8.5, 1.0, 20, False, TEXT)
text(s, "Objetivos específicos", 0.95, 4.05, 3.2, 0.4, 20, True, BLUE)
bullets(s, ["Validar los datos numéricos ingresados.",
            "Calcular el recargo escalonado.",
            "Generar un reporte consolidado."], 1.0, 4.55, 7.2, 1.7, 19)
for yy in [2.45, 3.65, 4.85]:
    rect(s, 9.0, yy, 1.05, 0.85, GREEN, True)
    text(s, "✓", 9.12, yy + 0.05, 0.8, 0.65, 34, True, WHITE, PP_ALIGN.CENTER)

# 5. Alternativas
s = base_slide("ALTERNATIVAS EVALUADAS", 5)
items = [
    ("Software empresarial", "Licenciamiento elevado y curva de aprendizaje."),
    ("Plantillas en Excel", "Fórmulas frágiles y riesgo de modificaciones accidentales."),
    ("Programa en Python", "Sin costo de licencia, simple, mantenible e iterativo.")
]
for i, (head, body) in enumerate(items):
    y = 2.25 + i * 1.35
    fill = BLUE_LIGHT if i < 2 else RGBColor(218, 246, 235)
    rect(s, 0.95, y, 8.7, 1.05, fill, True)
    text(s, head, 1.2, y + 0.14, 2.8, 0.3, 18, True, NAVY)
    text(s, body, 3.55, y + 0.14, 5.7, 0.55, 16, False, MUTED)
    text(s, "✓" if i == 2 else "×", 9.85, y + 0.05, 0.55, 0.75,
         31, True, GREEN if i == 2 else ORANGE, PP_ALIGN.CENTER)

# 6. Arquitectura
s = base_slide("SOLUCIÓN ELEGIDA: ARQUITECTURA", 6)
text(s, "Programa de consola organizado en 5 bloques lógicos:", 0.95, 2.25,
     7.9, 0.45, 19, False, MUTED)
labels = ["Ingreso y bucle", "Validación", "Motor de cálculo", "Arreglo", "Reporte final"]
coords = [(1.0, 3.15), (3.15, 3.15), (5.3, 3.15), (7.45, 3.15), (5.3, 4.65)]
for i, (lab, (x, y)) in enumerate(zip(labels, coords), 1):
    flow_box(s, f"{i}. {lab}", x, y, 1.85, 0.85, WHITE)
for x1, y1, x2, y2 in [(2.85, 3.55, 3.15, 3.55),
                       (5.0, 3.55, 5.3, 3.55),
                       (7.15, 3.55, 7.45, 3.55),
                       (8.35, 4.0, 6.3, 4.65)]:
    line = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = BLUE
    line.line.width = Pt(2.5)

# 7. Detalle del proceso
s = base_slide("DETALLE DEL PROCESO", 7)
flow_box(s, "ENTRADAS\nCantidad, nombre, sueldo/hora y horas extra",
         0.9, 2.7, 2.8, 1.35, BLUE_LIGHT)
flow_box(s, "PROCESO\nValidar, calcular y almacenar",
         4.25, 2.7, 2.8, 1.35, WHITE)
flow_box(s, "SALIDAS\nPago individual y total a depositar",
         7.6, 2.7, 2.8, 1.35, RGBColor(218, 246, 235))
for x1, x2 in [(3.7, 4.25), (7.05, 7.6)]:
    line = s.shapes.add_connector(1, Inches(x1), Inches(3.38), Inches(x2), Inches(3.38))
    line.line.color.rgb = BLUE
    line.line.width = Pt(3)
text(s, "El mismo flujo se repite automáticamente por cada técnico ingresado.",
     1.35, 4.75, 8.6, 0.55, 19, True, NAVY, PP_ALIGN.CENTER)

# 8. Algoritmo
s = base_slide("ALGORITMO (PSEUDOCÓDIGO)", 8)
code = (
    "1. Leer cantidad de técnicos\n"
    "2. Repetir por cada técnico\n"
    "3. Validar nombre, sueldo y horas\n"
    "4. Si horas_extra <= 2:\n"
    "      pago = horas_extra × sueldo_hora × 1.25\n"
    "5. Si no:\n"
    "      pago = 2 × sueldo_hora × 1.25\n"
    "             + (horas_extra - 2) × sueldo_hora × 1.35\n"
    "6. Guardar resultado y generar reporte"
)
rect(s, 0.95, 2.2, 8.9, 3.8, RGBColor(24, 27, 38), True)
text(s, code, 1.25, 2.45, 8.3, 3.25, 17, False,
     RGBColor(225, 232, 255), font="Courier New")
badge(s, "EJEMPLO: 3 h a S/15.50 = S/59.68", 1.5, 6.2, 5.7, GREEN)

# 9. Herramientas
s = base_slide("HERRAMIENTAS Y CONTROL DE VERSIONES", 9)
tools = [
    ("PYTHON 3", "Lógica, validaciones y cálculo monetario."),
    ("TRELLO", "Gestión del trabajo mediante Kanban e historias de usuario."),
    ("GIT / GITHUB", "Historial de versiones: estructura, lógica, validación y entrega final.")
]
for i, (head, body) in enumerate(tools):
    y = 2.25 + i * 1.35
    rect(s, 0.95, y, 8.9, 1.05, WHITE, True)
    badge(s, head, 1.15, y + 0.25, 2.1,
          BLUE if i != 1 else RGBColor(0, 121, 191))
    text(s, body, 3.55, y + 0.24, 5.9, 0.5, 17, False, TEXT)

# 10. Demostración
s = base_slide("DEMOSTRACIÓN EN VIVO", 10)
text(s, "A continuación, ejecutaremos el programa con tres casos de prueba.",
     1.0, 2.25, 8.5, 0.55, 21, True, NAVY)
for i, (lab, val) in enumerate([
    ("Caso 1", "3 h extra: supera el límite"),
    ("Caso 2", "1.5 h extra: no supera el límite"),
    ("Caso 3", "5 h extra: aplica ambos recargos")
]):
    y = 3.15 + i * 0.85
    badge(s, lab, 1.1, y, 1.2, BLUE)
    text(s, val, 2.55, y + 0.07, 5.8, 0.35, 18, False, TEXT)
rect(s, 8.8, 2.45, 1.15, 3.45, RGBColor(24, 27, 38), True)
text(s, ">_", 8.94, 3.55, 0.9, 0.8, 31, True, GREEN, PP_ALIGN.CENTER)

# 11. Resultados
s = base_slide("RESULTADOS", 11)
bullets(s, [
    "Técnico con 3 horas extra y sueldo de S/15.50 por hora: pago de S/59.68.",
    "El sistema calcula automáticamente y aplica redondeo monetario.",
    "El reporte final muestra cada pago y el total formateado a dos decimales."
], 0.95, 2.45, 7.6, 2.8, 20)
rect(s, 8.75, 2.2, 1.35, 3.5, RGBColor(24, 27, 38), True)
text(s, "S/\n59.68", 8.88, 3.0, 1.1, 1.3, 30, True, GREEN,
     PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

# 12. Conclusiones
s = base_slide("CONCLUSIONES Y RECOMENDACIONES", 12)
text(s, "Conclusiones", 0.95, 2.25, 2.4, 0.4, 20, True, BLUE)
bullets(s, [
    "Se automatizó el cálculo del recargo escalonado, reduciendo errores manuales.",
    "Los arreglos permiten procesar varios técnicos sin reescribir el código."
], 0.95, 2.75, 8.7, 1.45, 18)
text(s, "Recomendaciones", 0.95, 4.45, 2.8, 0.4, 20, True, BLUE)
bullets(s, [
    "Exportar el reporte a Excel o PDF.",
    "Conectar el programa a una base de datos para conservar el historial de pagos."
], 0.95, 4.9, 8.7, 1.25, 18)

# 13. Cierre
s = prs.slides.add_slide(blank)
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 0, 0, 3.0, 7.5, BLUE_DARK)
rect(s, 2.6, 0, 0.35, 7.5, BLUE)
text(s, "GRACIAS", 3.5, 1.45, 6.8, 0.8, 40, True, WHITE, PP_ALIGN.CENTER)
text(s, "Grupo 4", 3.5, 2.35, 6.8, 0.45, 21, False, BLUE_LIGHT, PP_ALIGN.CENTER)
for i, name in enumerate([
    "Christian Anthony Pozo Mejía",
    "Piero Alexander Linares Martínez",
    "David Laura Gonzales"
]):
    text(s, name, 3.6, 3.35 + i * 0.55, 6.6, 0.35, 18, False,
         WHITE, PP_ALIGN.CENTER)
text(s, "LIQUIDADOR DE HORAS EXTRAS AUTOMATIZADO", 3.2, 5.7, 7.4,
     0.45, 16, True, BLUE_LIGHT, PP_ALIGN.CENTER)

prs.save(OUT)
print(f"Presentación generada: {OUT}")
